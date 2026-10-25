import re
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

NAVY_HEX = "1E3A8A"
NAVY_RGB = RGBColor(0x1E, 0x3A, 0x8A)
DARK_RGB = RGBColor(0x0F, 0x17, 0x2A)
MUTED_RGB = RGBColor(0x47, 0x55, 0x69)


def _set_cell_margins(cell, top=20, bottom=20, left=0, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _add_bottom_border(paragraph, color_hex=NAVY_HEX, size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_top_border(paragraph, color_hex="CBD5E1", size="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), size)
    top.set(qn("w:space"), "2")
    top.set(qn("w:color"), color_hex)
    pBdr.append(top)
    pPr.append(pBdr)


def _append_html_runs(paragraph, html_text: str, font_size_pt: float = 9.0, color: RGBColor = DARK_RGB):
    text = html_text.replace("&amp;", "&").replace("&bull;", "•").replace("&nbsp;", " ")
    parts = re.split(r"(<strong>.*?</strong>)", text)
    for part in parts:
        if not part:
            continue
        is_bold = False
        clean_part = part
        if part.startswith("<strong>") and part.endswith("</strong>"):
            is_bold = True
            clean_part = part[8:-9]
        run = paragraph.add_run(clean_part)
        run.bold = is_bold
        run.font.name = "Segoe UI"
        run.font.size = Pt(font_size_pt)
        run.font.color.rgb = color


def _add_section_heading(doc: Document, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.name = "Segoe UI"
    run.font.size = Pt(9.5)
    run.font.color.rgb = NAVY_RGB
    _add_bottom_border(p, color_hex="CBD5E1", size="6")


def build_nancy_cv_docx(
    profile: Dict[str, Any],
    output_path: Path,
    custom_headline: Optional[str] = None,
) -> Path:
    p = profile["personal_info"]
    headline = custom_headline or p["role_headline_de"]
    photo_file = DATA_DIR / p["photo_path"]
    sig_file = DATA_DIR / p["signature_path"]
    today_str = datetime.now().strftime("%d.%m.%Y")

    doc = Document()

    # Configure A4 page & tight 1-page margins
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(9)
    section.bottom_margin = Mm(8)
    section.left_margin = Mm(12)
    section.right_margin = Mm(12)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Segoe UI"
    normal_style.font.size = Pt(9)

    # 1. Header Table (Left: Name & Contact, Right: Portrait Photo)
    hdr_table = doc.add_table(rows=1, cols=2)
    hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_table.autofit = False
    hdr_table.columns[0].width = Mm(154)
    hdr_table.columns[1].width = Mm(32)

    left_cell = hdr_table.cell(0, 0)
    right_cell = hdr_table.cell(0, 1)
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p_name = left_cell.paragraphs[0]
    p_name.paragraph_format.space_after = Pt(1)
    r_name = p_name.add_run(p["full_name"])
    r_name.bold = True
    r_name.font.name = "Segoe UI"
    r_name.font.size = Pt(17.5)
    r_name.font.color.rgb = DARK_RGB

    p_hl = left_cell.add_paragraph()
    p_hl.paragraph_format.space_after = Pt(2)
    r_hl = p_hl.add_run(p.get("motto_keywords", headline))
    r_hl.bold = True
    r_hl.font.name = "Segoe UI"
    r_hl.font.size = Pt(9.8)
    r_hl.font.color.rgb = NAVY_RGB

    if p.get("about_me_de"):
        p_abt = left_cell.add_paragraph()
        p_abt.paragraph_format.space_after = Pt(3)
        r_abt = p_abt.add_run(p["about_me_de"])
        r_abt.font.name = "Segoe UI"
        r_abt.font.size = Pt(8.3)
        r_abt.font.color.rgb = MUTED_RGB

    p_c1 = left_cell.add_paragraph()
    p_c1.paragraph_format.space_after = Pt(1.5)
    _append_html_runs(
        p_c1,
        f"<strong>Adresse:</strong> {p['address']}    |    <strong>Mobil:</strong> {p['phone']}",
        font_size_pt=8.5,
        color=MUTED_RGB,
    )

    p_c2 = left_cell.add_paragraph()
    p_c2.paragraph_format.space_after = Pt(2)
    _append_html_runs(
        p_c2,
        f"<strong>E-Mail:</strong> {p['email']}    |    <strong>LinkedIn:</strong> {p.get('linkedin_label', 'LinkedIn Profile')}",
        font_size_pt=8.5,
        color=MUTED_RGB,
    )

    if photo_file.exists():
        p_img = right_cell.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_img.paragraph_format.space_after = Pt(0)
        p_img.add_run().add_picture(str(photo_file), width=Mm(26))

    # Divider line under header
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(2.5)
    _add_bottom_border(p_div, color_hex=NAVY_HEX, size="14")

    # 2. Berufserfahrung Section
    _add_section_heading(doc, "Berufserfahrung")
    exp_table = doc.add_table(rows=len(profile["experience_de"]), cols=2)
    exp_table.autofit = False
    for idx, exp in enumerate(profile["experience_de"]):
        row = exp_table.rows[idx]
        c_date = row.cells[0]
        c_body = row.cells[1]
        c_date.width = Mm(28)
        c_body.width = Mm(158)
        _set_cell_margins(c_date, top=18, bottom=18, left=0, right=50)
        _set_cell_margins(c_body, top=18, bottom=18, left=0, right=0)

        p_dt = c_date.paragraphs[0]
        p_dt.paragraph_format.space_after = Pt(0)
        r_dt = p_dt.add_run(exp["dates"])
        r_dt.bold = True
        r_dt.font.name = "Segoe UI"
        r_dt.font.size = Pt(8.4)
        r_dt.font.color.rgb = MUTED_RGB

        p_title = c_body.paragraphs[0]
        p_title.paragraph_format.space_after = Pt(1)
        _append_html_runs(
            p_title,
            f"<strong>{exp['title']}</strong> | <strong>{exp['company']}</strong>, {exp['location']}",
            font_size_pt=8.9,
            color=DARK_RGB,
        )

        for b in exp.get("bullets", []):
            p_b = c_body.add_paragraph()
            p_b.paragraph_format.left_indent = Mm(3.5)
            p_b.paragraph_format.first_line_indent = Mm(-3.5)
            p_b.paragraph_format.space_after = Pt(1)
            _append_html_runs(p_b, f"•  {b}", font_size_pt=8.4, color=DARK_RGB)

    # 3. Bildungsweg Section
    _add_section_heading(doc, "Bildungsweg")
    edu_table = doc.add_table(rows=len(profile["education_de"]), cols=2)
    edu_table.autofit = False
    for idx, edu in enumerate(profile["education_de"]):
        row = edu_table.rows[idx]
        c_date = row.cells[0]
        c_body = row.cells[1]
        c_date.width = Mm(28)
        c_body.width = Mm(158)
        _set_cell_margins(c_date, top=16, bottom=16, left=0, right=50)
        _set_cell_margins(c_body, top=16, bottom=16, left=0, right=0)

        p_dt = c_date.paragraphs[0]
        p_dt.paragraph_format.space_after = Pt(0)
        r_dt = p_dt.add_run(edu["dates"])
        r_dt.bold = True
        r_dt.font.name = "Segoe UI"
        r_dt.font.size = Pt(8.4)
        r_dt.font.color.rgb = MUTED_RGB

        p_deg = c_body.paragraphs[0]
        p_deg.paragraph_format.space_after = Pt(1)
        _append_html_runs(
            p_deg,
            f"<strong>{edu['degree']}</strong> | {edu['institution']}",
            font_size_pt=8.8,
            color=DARK_RGB,
        )
        if edu.get("details"):
            p_det = c_body.add_paragraph()
            p_det.paragraph_format.space_after = Pt(1)
            _append_html_runs(p_det, edu["details"], font_size_pt=8.2, color=MUTED_RGB)

    # 4. Fachkompetenzen & IT-Kenntnisse
    _add_section_heading(doc, "Fachkompetenzen & IT-Kenntnisse")
    lang_str = "  •  ".join(f"<strong>{l['language']}:</strong> {l['level']}" for l in profile["languages_de"])

    skills_lines = [
        f"<strong>Office & Organisation:</strong> {profile['skills_de']['office_management']}",
        f"<strong>Software & IT:</strong> {profile['skills_de']['it_software']}",
        f"<strong>Film & Audio:</strong> {profile['skills_de'].get('film_audio', 'DaVinci Resolve · Wondershare Filmora · Adobe Audition CC 2019')}",
        f"<strong>Sprachen:</strong> {lang_str}",
        f"<strong>Nachweise & Mobilität:</strong> {profile['skills_de']['certificates']}",
    ]
    for s_line in skills_lines:
        ps = doc.add_paragraph()
        ps.paragraph_format.space_after = Pt(2)
        _append_html_runs(ps, s_line, font_size_pt=8.45, color=DARK_RGB)

    # 5. Footer Signature Table
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.autofit = False
    sig_table.columns[0].width = Mm(120)
    sig_table.columns[1].width = Mm(66)

    c_loc = sig_table.cell(0, 0)
    c_sig = sig_table.cell(0, 1)
    c_loc.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    c_sig.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM

    p_loc = c_loc.paragraphs[0]
    p_loc.paragraph_format.space_before = Pt(8)
    p_loc.paragraph_format.space_after = Pt(0)
    r_loc = p_loc.add_run(f"Berlin, den {today_str}")
    r_loc.font.name = "Segoe UI"
    r_loc.font.size = Pt(8.6)
    r_loc.font.color.rgb = MUTED_RGB

    p_sig = c_sig.paragraphs[0]
    p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig.paragraph_format.space_before = Pt(4)
    p_sig.paragraph_format.space_after = Pt(1)
    if sig_file.exists():
        p_sig.add_run().add_picture(str(sig_file), height=Mm(9.5))

    p_sig_name = c_sig.add_paragraph()
    p_sig_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig_name.paragraph_format.space_after = Pt(0)
    _add_top_border(p_sig_name, color_hex="94A3B8", size="6")
    r_sn = p_sig_name.add_run(p["full_name"])
    r_sn.font.name = "Segoe UI"
    r_sn.font.size = Pt(8.4)
    r_sn.font.color.rgb = MUTED_RGB

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
